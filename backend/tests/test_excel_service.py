from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errors import AppError
from app.services.excel_service import create_workbook, parse_workbook


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_excel_round_trip_has_expected_headers_and_values():
    data = create_workbook([("问题一", "答案一"), ("Question", "Answer")])
    workbook = load_workbook(BytesIO(data), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert rows == [("问题", "固定答案"), ("问题一", "答案一"), ("Question", "Answer")]
    parsed = parse_workbook(data, 100)
    assert [(item.question, item.answer) for item in parsed] == [
        ("问题一", "答案一"),
        ("Question", "Answer"),
    ]


def test_excel_rejects_wrong_headers_and_empty_cells():
    with pytest.raises(AppError, match="第一行"):
        parse_workbook(workbook_bytes([("Q", "A"), ("x", "y")]), 10)
    with pytest.raises(AppError, match="第 2 行"):
        parse_workbook(workbook_bytes([("问题", "固定答案"), ("x", "")]), 10)
