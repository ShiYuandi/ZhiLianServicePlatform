from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.errors import AppError
from app.schemas.qa import QaItemCreate


HEADERS = ("问题", "固定答案")


def parse_workbook(data: bytes, max_rows: int) -> list[QaItemCreate]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise AppError("INVALID_EXCEL", "无法读取 Excel 文件", 422) from exc
    try:
        sheet = workbook.active
        first_row = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        values = tuple(
            "" if value is None else str(value).strip() for value in first_row
        )
        if values != HEADERS:
            raise AppError(
                "INVALID_EXCEL_HEADERS", "第一行必须且只能是“问题、固定答案”", 422
            )
        items: list[QaItemCreate] = []
        for excel_row, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if len(row) > 2 and any(value not in (None, "") for value in row[2:]):
                raise AppError(
                    "INVALID_EXCEL_COLUMNS", f"第 {excel_row} 行包含额外列", 422
                )
            question = "" if not row or row[0] is None else str(row[0]).strip()
            answer = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
            if not question and not answer:
                continue
            if not question or not answer:
                raise AppError(
                    "INVALID_EXCEL_ROW", f"第 {excel_row} 行问题或固定答案为空", 422
                )
            try:
                items.append(
                    QaItemCreate(
                        question=question, answer=answer, sort_order=len(items)
                    )
                )
            except ValueError as exc:
                raise AppError(
                    "INVALID_EXCEL_ROW", f"第 {excel_row} 行内容不符合长度要求", 422
                ) from exc
            if len(items) > max_rows:
                raise AppError(
                    "EXCEL_TOO_MANY_ROWS", f"Excel 最多允许 {max_rows} 条问答", 413
                )
        if not items:
            raise AppError("EMPTY_EXCEL", "Excel 中没有可导入的问答", 422)
        return items
    finally:
        workbook.close()


def create_workbook(rows: list[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "问答数据"
    sheet.append(list(HEADERS))
    for question, answer in rows:
        sheet.append([question, answer])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{max(1, sheet.max_row)}"
    widths = [40, 80]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
